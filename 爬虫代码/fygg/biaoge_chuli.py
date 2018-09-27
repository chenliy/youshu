#!/usr/bin/env python
# -*- encoding: utf-8 -*-
# Created on 2018/7/27 9:24
# Project: 
# @Author: ZQJ
# @Email : zihe@yscredit.com


import pandas as pd
import docx
from docx import Document
path = r'C:\Users\ll\Downloads\15127211182030501570.docx'
document = Document(path)
tables = document.tables

for table in tables:
    len_row = len(table.rows)#行数
    print(len_row)
    len_loc = len(table.columns)#列数
    print(len_loc)
    json_result = {}
    print(table.cell(0,0).text)
    json_result['title'] = table.cell(0,0).text
    json_result['telephone'] = table.cell(1,0).text
    print(table.cell(1,0).text)
    json_result['content_list'] = []

    a = []
    for row in table.rows:
        content = []
        for i in range(len_loc):
            content.append(row.cells[i].text)
        a.append(content)

    print(a[1])
    print(a[2])
    print(a[3])
    print(a[4])
    for i in range(4,len(a)):
        content_one = {}
        content_one['序号'] = a[i][0]
        content_one['姓名'] = a[i][1]
        content_one['性别'] = a[i][2]
        content_one['出生日期'] = a[i][3]
        content_one['籍贯'] = a[i][4]
        content_one['罪名'] = a[i][5]
        content_one['原判'] = a[i][6]
        content_one['起日'] = a[i][7]
        content_one['现止日'] = a[i][8]
        content_one['历次减刑情况'] = a[i][9]
        content_one['提请减刑的主要理由'] = a[i][10]
        content_one['刑罚执行机关提请减刑意见'] = a[i][11]
        content_one['备注'] = a[i][12]
        json_result['content_list'].append(content_one)
    print(json_result)



url = 'http://kmzy.ynfy.gov.cn/u/cms/www/201712/19105216hul5.doc'
import requests

r = requests.get(url).content
f = open(r,'rb')
print(f)

import subprocess
subprocess.check_output()

